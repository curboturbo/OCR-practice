import os
import cv2
from scan_boxes import Recognize
import json
from openpyxl import Workbook

"""СОЗДАНИЕ ДАТАСЕТА ДЛЯ BERT"""
rows = 1
wb = Workbook()
ws = wb.active


def find(a: int, mes: str) -> int:
    a = a - 21
    last = mes[0]
    cnt = 0
    for i in range(len(mes)):
        if mes[i] == ')':
            cnt += 1
        if cnt == a:
            return mes[i - 3]


def write_to_xl(answer, mes) -> None:
    global rows
    global ws
    if answer:
        for k, v in answer.items():
            ws[f'A{rows}'] = k[-2] + k[-1]
            ws[f'B{rows}'] = v
            ws[f'C{rows}'] = find(a=int(k[-2] + k[-1]), mes=mes)
            rows += 1


def make_list() -> None:
    PATH_301 = r"C:\Users\Tema\Desktop\301"
    PATH_JSON = r"C:\Users\Tema\Desktop\school"
    content = os.listdir(PATH_JSON)
    content_301 = os.listdir(PATH_301)
    data = {}
    jj = 0
    for i in content_301:
        if jj == 1000:
            break
        cur = i.replace('-', '').split('.')[0].split('_')[0]
        path_find = rf"{cur}_02__1_res.txt.webRes"

        if path_find in content:
            if i.split('.')[-1] == 'json':
                aa = os.path.join(PATH_301, i)  # json _301
                with open(f'{aa}', 'r', encoding='utf-8') as f:
                    data = json.load(f)['mask']

            if i.split('.')[-1] == 'png':
                path = r'C:\Users\Tema\Desktop\leetcode\runs\segment\train\weights\best.pt'
                ab = os.path.join(PATH_301, i)
                js = os.path.join(PATH_JSON, path_find)
                obj = Recognize(json_file_path=js, photo=ab, model=path, sensitivity=0.3).main_function()
                write_to_xl(obj, data)
                jj += 1
                input()


make_list()
wb.save("bert.xlsx")
print('all was writed')